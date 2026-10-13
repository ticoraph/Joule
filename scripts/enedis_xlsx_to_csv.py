"""Convert Enedis .xlsx exports (from scripts/enedis_download.py) into CSVs.

The xlsx files have a metadata sheet ("Page d'accueil") plus a data sheet
whose table doesn't start at A1 and whose header row position varies. This
locates the header row by matching known column labels, wherever they are,
then writes a CSV with column names matching data/processed/*.csv so the
output is compatible with scripts/import_csv.py.

Usage (from repo root):
    python scripts/enedis_xlsx_to_csv.py [--raw-dir data/raw]

Picks the most recent energy_*.xlsx, index_*.xlsx and powermax_*.xlsx in
--raw-dir and writes a sibling .csv next to each.
"""

import argparse
import csv
import logging
from datetime import datetime
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RAW_DIR = REPO_ROOT / "data" / "raw"
LOG_DIR = REPO_ROOT / "logs"

LOG_DIR.mkdir(parents=True, exist_ok=True)
logger = logging.getLogger("enedis_xlsx_to_csv")
logger.setLevel(logging.INFO)
_handler = logging.FileHandler(LOG_DIR / "enedis_xlsx_to_csv.log", encoding="utf-8")
_handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
logger.addHandler(_handler)
logger.addHandler(logging.StreamHandler())

# output column name -> exact header label to find in the xlsx
EXPORTS = {
    "energy": {"Date": "Date", "Conso": "Valeur (en kWh)"},
    "index": {
        "Date": "Date du télé-relevé",
        "Index_Total": "Index totalisateur (en kWh)",
        "Index_HP": "Heures Pleines (en kWh)",
        "Index_HC": "Heures Creuses (en kWh)",
    },
    "powermax": {"Date": "Date et heure", "PMax": "Valeur (en kVA)"},
}

# export prefix -> the Date column's source format in the xlsx (Enedis exports
# this as text, not a real date cell)
SOURCE_DATE_FORMATS = {
    "energy": "%d/%m/%Y",
    "index": "%d/%m/%Y",
    "powermax": "%d/%m/%Y %H:%M",
}

# export prefix -> name of an extra output column holding the HH:MM time part
# split out of the source Date (only powermax's Date carries a time of day)
HOUR_COLUMNS = {"powermax": "PMax_Hour"}


def find_header_row(ws, labels: set) -> tuple[dict, int]:
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        found = {
            str(cell.value).strip(): cell.column
            for cell in row
            if cell.value is not None and str(cell.value).strip() in labels
        }
        if labels <= found.keys():
            return found, row[0].row
    raise ValueError(f"En-tete introuvable pour {labels}")


def convert(xlsx_path: Path, columns: dict, source_date_format: str, hour_column: str | None) -> Path:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    # Normally one data sheet, but Enedis splits the year's export across
    # several sheets (e.g. "Index_Cons_Calendriers_1"/"_2") when there's a
    # mid-year change affecting the export (observed in 2026 on a supplier
    # switch) - concatenate every non-summary sheet's rows in order.
    data_sheets = [s for s in wb.sheetnames if s != "Page d'accueil"]

    out_columns = list(columns.keys())
    if hour_column:
        out_columns.insert(out_columns.index("Date") + 1, hour_column)

    csv_path = xlsx_path.with_suffix(".csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(out_columns)
        for sheet_name in data_sheets:
            ws = wb[sheet_name]
            col_for_label, header_row = find_header_row(ws, set(columns.values()))
            col_idx = {out_name: col_for_label[label] for out_name, label in columns.items()}
            for r in range(header_row + 1, ws.max_row + 1):
                date_value = ws.cell(row=r, column=col_idx["Date"]).value
                if date_value is None:
                    continue
                dt = datetime.strptime(str(date_value).strip(), source_date_format)
                row = []
                for c in out_columns:
                    if c == "Date":
                        row.append(dt.strftime("%Y%m%d"))
                    elif c == hour_column:
                        row.append(dt.strftime("%H:%M"))
                    else:
                        row.append(ws.cell(row=r, column=col_idx[c]).value)
                writer.writerow(row)

    return csv_path


def latest_file(raw_dir: Path, prefix: str) -> Path:
    matches = sorted(raw_dir.glob(f"{prefix}_*.xlsx"))
    if not matches:
        raise SystemExit(f"Aucun fichier {prefix}_*.xlsx dans {raw_dir}")
    return matches[-1]


def main(raw_dir: Path) -> None:
    for prefix, columns in EXPORTS.items():
        xlsx_path = latest_file(raw_dir, prefix)
        csv_path = convert(xlsx_path, columns, SOURCE_DATE_FORMATS[prefix], HOUR_COLUMNS.get(prefix))
        logger.info(f"{xlsx_path.name} -> {csv_path.name}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW_DIR)
    args = parser.parse_args()
    main(args.raw_dir)
